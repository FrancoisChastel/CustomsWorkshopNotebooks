"""The Lock sheet of a session workbook: the three parts of every sprint's lock as yellow cells.

One writer for two uses: the shipped template carries the empty form (a group fills it when the
site is down and hands it to the facilitator), and the hub's export fills the same form from the
group's locks (ADD-J J7). The fields come from hosted/hub/sprints/locks.json, the one schema the
hub validates against and the site renders; a copy travels inside this package as lock_fields.json.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

NAVY, INK, MUTED = "0E2F5A", "3C4A5E", "6A7686"
LEVER = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor=NAVY)
SESSION_TITLES = {"1b": "Breakout 1B · Targeting", "2b": "Breakout 2B · Trade and the border", "4b": "Breakout 4B · Informality"}
PARTS = (("says", "What the data says"), ("decide", "What we decide"), ("signs", "Who signs and what it costs"))
CANDIDATES = [Path(__file__).resolve().parents[2] / "hosted" / "hub" / "sprints" / "locks.json", Path("/opt/hub/sprints/locks.json"), Path(__file__).resolve().parent / "lock_fields.json"]


def lock_fields(path: str | Path | None = None) -> dict[str, Any]:
    """The lock schema, from the repository, the hub image or the packaged copy, whichever exists first."""
    for candidate in ([Path(path)] if path else []) + CANDIDATES:
        if candidate.exists():
            return json.loads(candidate.read_text(encoding="utf-8"))
    raise FileNotFoundError("locks.json: no lock schema found")


def sprints_of(session: str, fields: dict[str, Any] | None = None) -> list[tuple[str, dict[str, Any]]]:
    fields = fields or lock_fields()
    return sorted(((key, spec) for key, spec in fields.items() if key.startswith(f"{session.lower()}-")), key=lambda item: int(item[0].split("-")[1]))


def _font(bold: bool = False, color: str = INK, size: int = 10, italic: bool = False) -> Font:
    return Font(name="Arial", bold=bold, color=color, size=size, italic=italic)


def _lever(cell, value: Any = None) -> None:
    cell.fill = LEVER
    cell.font = _font(bold=True, color=NAVY)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if value is not None and value != "":
        cell.value = value


def _text(value: Any) -> Any:
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return value


def lock_sheet(ws: Worksheet, session: str, locks: dict[str, Any] | None = None, fields: dict[str, Any] | None = None, group: str = "") -> dict[str, tuple[int, int]]:
    """Write the Lock sheet for `session`; `locks` maps a sprint key to its lock entry (says, decide, signs).

    Returns the first cell (row, column) of every sprint block, for the reference sheets to point at.
    """
    locks = locks or {}
    ws.title = "Lock"
    ws["A1"] = f"{SESSION_TITLES.get(session.lower(), session)} · the locks"
    ws["A1"].font = _font(bold=True, color=NAVY, size=14)
    ws["A2"] = ("Every sprint locks three parts: what the data says, what we decide, who signs and what it costs. Yellow cells are yours. "
                "If the site is down, fill this sheet and hand it to the facilitator." + (f" Group: {group}." if group else ""))
    ws["A2"].font = _font(color=INK)
    ws.column_dimensions["A"].width = 46
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 22
    row = 4
    anchors: dict[str, tuple[int, int]] = {}
    for key, spec in sprints_of(session, fields):
        entry = locks.get(key) or {}
        anchors[key] = (row, 1)
        ws.cell(row=row, column=1, value=f"Sprint {key.split('-')[1]} · {spec['title']}").font = _font(bold=True, color="FFFFFF", size=11)
        for col in range(1, 9):
            ws.cell(row=row, column=col).fill = HEAD
        if entry.get("at"):
            ws.cell(row=row, column=8, value=f"locked {str(entry['at'])[:16].replace('T', ' ')}").font = _font(color="FFFFFF", italic=True)
        row += 1
        for part, label in PARTS:
            ws.cell(row=row, column=1, value=label).font = _font(bold=True, color=NAVY)
            row += 1
            if part == "says":
                ws.cell(row=row, column=1, value=spec["says"]).font = _font()
                ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
                _lever(ws.cell(row=row, column=2), _text(entry.get("says")))
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
                ws.row_dimensions[row].height = 48
                row += 1
                continue
            decided = entry.get("decide") or {} if part == "decide" else entry.get("signs") or {}
            for field in spec[part]:
                if field.get("kind") == "table":
                    ws.cell(row=row, column=1, value=field["label"] + (" (required)" if field.get("required") else "")).font = _font()
                    row += 1
                    ws.cell(row=row, column=1, value="row").font = _font(bold=True, color=MUTED)
                    for c, column in enumerate(field["columns"], 2):
                        cell = ws.cell(row=row, column=c, value=column["label"] + (" · " + " / ".join(column["options"]) if column.get("options") else ""))
                        cell.font = _font(bold=True, color=MUTED)
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                    row += 1
                    values = decided.get(field["key"]) or {}
                    for row_id in field["rows"]:
                        ws.cell(row=row, column=1, value=row_id).font = _font()
                        given = values.get(row_id) if isinstance(values, dict) else None
                        for c, column in enumerate(field["columns"], 2):
                            _lever(ws.cell(row=row, column=c), _text((given or {}).get(column["key"])) if isinstance(given, dict) else None)
                        row += 1
                else:
                    label = field["label"] + ("" if field.get("required") else " (optional)")
                    ws.cell(row=row, column=1, value=label).font = _font()
                    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
                    _lever(ws.cell(row=row, column=2), _text(decided.get(field["key"])))
                    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
                    row += 1
        row += 1
    ws.freeze_panes = "A4"
    return anchors
